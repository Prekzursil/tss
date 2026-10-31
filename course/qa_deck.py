"""QA deck.pptx: estimated text overflow, shapes outside slide, leftover placeholders.

Usage:
    python qa_deck.py [path-to-deck.pptx]

Exit code is 0 when no issues are found, 1 otherwise (CI-friendly). The deck
path defaults to ``deck.pptx`` next to this script.
"""

from __future__ import annotations

import math
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

SW, SH = 13.333, 7.5
# rough advance width factor (fraction of font size) per font
WFAC = {"Segoe UI": 0.50, "Consolas": 0.55}

PLACEHOLDER_MARKERS = ("todo", "lorem", "xxx", "[insert", "click to add")


def _paragraph_est_height(p, w: float) -> float:
    """Estimated rendered height (inches) of one text paragraph in a `w`-inch frame."""
    txt = "".join(r.text for r in p.runs)
    if not txt:
        return 0.14
    size = max((r.font.size.pt if r.font.size else 14) for r in p.runs)
    font = next((r.font.name for r in p.runs if r.font.name), "Segoe UI")
    fac = WFAC.get(font, 0.5)
    cpl = max(int(w / (size * fac / 72)), 1) if w > 0 else 1
    lines = max(math.ceil(len(txt) / cpl), txt.count("\n") + 1)
    ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
    est = lines * size * 1.22 * ls / 72
    if p.space_after:
        est += p.space_after.pt / 72
    return est


def walk(shapes, sidx: int, issues: list[str]) -> None:
    for sp in shapes:
        if sp.shape_type == 6:
            walk(sp.shapes, sidx, issues)
            continue
        x = Emu(sp.left).inches if sp.left is not None else 0
        y = Emu(sp.top).inches if sp.top is not None else 0
        w = Emu(sp.width).inches if sp.width is not None else 0
        h = Emu(sp.height).inches if sp.height is not None else 0
        if x < -0.01 or y < -0.01 or x + w > SW + 0.01 or y + h > SH + 0.01:
            issues.append(
                f"slide {sidx}: shape '{sp.shape_type}' out of slide "
                f"({x:.2f},{y:.2f},{w:.2f},{h:.2f})"
            )
        if not getattr(sp, "has_text_frame", False) or not sp.has_text_frame:
            continue
        tf = sp.text_frame
        if not tf.text.strip():
            continue
        est_h = 0.0
        text_all: list[str] = []
        for p in tf.paragraphs:
            text_all.append("".join(r.text for r in p.runs))
            est_h += _paragraph_est_height(p, w)
        joined = " ".join(text_all).lower()
        for bad in PLACEHOLDER_MARKERS:
            if bad in joined:
                issues.append(f"slide {sidx}: placeholder text '{bad}' in {text_all[:1]}")
        if est_h > h * 1.18 + 0.06:
            issues.append(
                f"slide {sidx}: possible overflow ({est_h:.2f}in est vs {h:.2f}in box) "
                f"text={text_all[:1]}"
            )


def audit_deck(path: str | Path) -> tuple[int, list[str]]:
    """Audit the deck at `path`; return (slide_count, issues)."""
    prs = Presentation(str(path))
    issues: list[str] = []
    for i, sl in enumerate(prs.slides, 1):
        walk(sl.shapes, i, issues)
    return len(prs.slides._sldIdLst), issues


def main(argv: list[str] | None = None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)
    path = Path(argv[0]) if argv else Path(__file__).resolve().parent / "deck.pptx"
    slides, issues = audit_deck(path)
    print(f"deck: {path.name} — slides: {slides}")
    if issues:
        print(f"{len(issues)} issue(s):")
        for it in issues:
            print(" -", it)
        return 1
    print("no issues found")
    return 0


if __name__ == "__main__":
    sys.exit(main())
