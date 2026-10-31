"""Tests for the deck QA audit — synthetic decks built with python-pptx.

Every audit finding class (out-of-slide shape, placeholder text, estimated
overflow) plus the clean-deck happy path is exercised against decks built
in-memory with the same library the audit reads.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu, Inches, Pt

COURSE_ROOT = Path(__file__).resolve().parent.parent
if str(COURSE_ROOT) not in sys.path:
    sys.path.insert(0, str(COURSE_ROOT))

import qa_deck  # noqa: E402


def _blank():
    # 16:9 default template; python-pptx starts with zero slides, so add one blank
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.slides.add_slide(prs.slide_layouts[6])  # 6 = blank layout
    return prs


def _box(prs: Presentation, *, x=0.5, y=0.5, w=4.0, h=1.0) -> None:  # type: ignore  # pyright: ignore[reportGeneralTypeIssues]
    slide = prs.slides[0]
    slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))


def _textbox_with(prs, text: str, *, size=18, x=0.5, y=0.5, w=4.0, h=1.0):
    slide = prs.slides[0]
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.text = text
    tf.paragraphs[0].runs[0].font.size = Pt(size)
    return box


def test_clean_deck_reports_no_issues(tmp_path):
    p = tmp_path / "clean.pptx"
    prs = _blank()
    _box(prs)  # one empty, in-bounds textbox
    prs.save(str(p))
    slides, issues = qa_deck.audit_deck(p)
    assert slides == 1
    assert issues == []


def test_shape_outside_slide_is_flagged(tmp_path):
    p = tmp_path / "out.pptx"
    prs = _blank()
    # width pushes the right edge past SW (13.333): x + w > SW + 0.01
    _textbox_with(prs, "edge", x=12.0, y=0.5, w=4.0, h=1.0)
    prs.save(str(p))
    _, issues = qa_deck.audit_deck(p)
    assert any("out of slide" in i for i in issues)


def test_shape_negative_origin_is_flagged(tmp_path):
    p = tmp_path / "neg.pptx"
    prs = _blank()
    _textbox_with(prs, "off", x=-0.5, y=0.5, w=2.0, h=1.0)
    prs.save(str(p))
    _, issues = qa_deck.audit_deck(p)
    assert any("out of slide" in i for i in issues)


@pytest.mark.parametrize("marker", ["TODO", "lorem", "xxx", "[insert", "click to add"])
def test_placeholder_markers_are_flagged(tmp_path, marker):
    p = tmp_path / "ph.pptx"
    prs = _blank()
    _textbox_with(prs, f"heading with {marker} inside", size=14, w=8.0, h=1.0)
    prs.save(str(p))
    _, issues = qa_deck.audit_deck(p)
    assert any("placeholder text" in i for i in issues)


def test_estimated_overflow_is_flagged(tmp_path):
    p = tmp_path / "ovf.pptx"
    prs = _blank()
    # ~200 chars at 28pt in a 5-inch-wide, 0.3-inch-tall box: many wrapped lines
    _textbox_with(prs, "word " * 40, size=28, w=5.0, h=0.3)
    prs.save(str(p))
    _, issues = qa_deck.audit_deck(p)
    assert any("possible overflow" in i for i in issues)


def test_small_text_in_tall_box_is_not_overflow(tmp_path):
    p = tmp_path / "ok.pptx"
    prs = _blank()
    _textbox_with(prs, "fits", size=12, w=6.0, h=2.0)
    prs.save(str(p))
    _, issues = qa_deck.audit_deck(p)
    assert issues == []


def test_grouped_shapes_are_walked(tmp_path, monkeypatch):
    # A group (shape_type 6) is recursed into; put an out-of-slide box inside one.
    p = tmp_path / "grp.pptx"
    prs = _blank()
    slide = prs.slides[0]
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Emu(0), Emu(0), Inches(2.0), Inches(1.0))
    box.left = Inches(13.0)  # inside the group, right edge beyond SW
    box.width = Inches(2.0)
    tf = box.text_frame
    tf.text = "g"
    prs.save(str(p))
    _, issues = qa_deck.audit_deck(p)
    assert any("out of slide" in i for i in issues)


def test_paragraph_est_height_empty_paragraph():
    prs = _blank()
    box = prs.slides[0].shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "x"
    # second, empty paragraph
    p2 = box.text_frame.add_paragraph()
    assert qa_deck._paragraph_est_height(p2, 4.0) == 0.14


def test_paragraph_est_height_space_after_counts():
    from pptx.util import Pt

    prs = _blank()
    box = prs.slides[0].shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "spaced"
    p = box.text_frame.paragraphs[0]
    p.space_after = Pt(12)
    with_space = qa_deck._paragraph_est_height(p, 4.0)
    p.space_after = None
    without_space = qa_deck._paragraph_est_height(p, 4.0)
    assert with_space == pytest.approx(without_space + 12 / 72)


def test_main_reads_argv_when_called_with_none(monkeypatch, tmp_path, capsys):
    p = tmp_path / "argv.pptx"
    prs = _blank()
    _box(prs)
    prs.save(str(p))
    monkeypatch.setattr(sys, "argv", ["qa_deck.py", str(p)])
    assert qa_deck.main(None) == 0
    assert "no issues found" in capsys.readouterr().out


def test_paragraph_est_height_width_zero_clamps_cpl():
    prs = _blank()
    box = prs.slides[0].shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "some copy"
    p = box.text_frame.paragraphs[0]
    # w=0 must not divide by zero: cpl clamps to 1 -> many lines
    est = qa_deck._paragraph_est_height(p, 0.0)
    assert est > 0


def test_main_reports_zero_exit_on_clean_deck(tmp_path, capsys):
    p = tmp_path / "m.pptx"
    prs = _blank()
    _box(prs)
    prs.save(str(p))
    assert qa_deck.main([str(p)]) == 0
    out = capsys.readouterr().out
    assert "no issues found" in out


def test_main_reports_nonzero_exit_on_flagged_deck(tmp_path, capsys):
    p = tmp_path / "m2.pptx"
    prs = _blank()
    _textbox_with(prs, "lorem ipsum draft", size=14, w=6.0, h=1.0)
    prs.save(str(p))
    assert qa_deck.main([str(p)]) == 1
    out = capsys.readouterr().out
    assert "issue(s)" in out


def test_main_defaults_to_deck_next_to_script(capsys):
    # Default path: <script dir>/deck.pptx — the committed 20-slide deck audits clean.
    assert qa_deck.main([]) == 0
    out = capsys.readouterr().out
    assert "no issues found" in out
